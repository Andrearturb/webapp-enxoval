"""Testes de integração para os endpoints de exportação.

Usam o cliente HTTP com banco real (catalogo_no_banco) + enxoval criado via API.
"""
import io

import pytest

RESPOSTAS_VALIDAS = {
    "municipio_codigo": 4106902,  # Curitiba
    "data_prevista": "2027-06-01",
    "dias_entre_lavagens": 2,
    "moradia": "apartamento",
    "tem_carro": True,
    "orcamento": "intermediario",
    "primeiro_filho": True,
}


@pytest.fixture
def enxoval_id(cliente, catalogo_no_banco):
    """Cria um enxoval e devolve o UUID como string."""
    resposta = cliente.post("/api/v1/enxovais", json=RESPOSTAS_VALIDAS)
    assert resposta.status_code == 201
    return resposta.json()["id"]


class TestExportarCsv:
    def test_retorna_200(self, cliente, enxoval_id):
        r = cliente.get(f"/api/v1/enxovais/{enxoval_id}/exportar.csv")
        assert r.status_code == 200

    def test_content_type_csv(self, cliente, enxoval_id):
        r = cliente.get(f"/api/v1/enxovais/{enxoval_id}/exportar.csv")
        assert "text/csv" in r.headers["content-type"]

    def test_content_disposition_attachment(self, cliente, enxoval_id):
        r = cliente.get(f"/api/v1/enxovais/{enxoval_id}/exportar.csv")
        assert "attachment" in r.headers["content-disposition"]
        assert ".csv" in r.headers["content-disposition"]

    def test_corpo_nao_vazio(self, cliente, enxoval_id):
        r = cliente.get(f"/api/v1/enxovais/{enxoval_id}/exportar.csv")
        assert len(r.content) > 0

    def test_404_para_enxoval_inexistente(self, cliente):
        r = cliente.get("/api/v1/enxovais/00000000-0000-0000-0000-000000000000/exportar.csv")
        assert r.status_code == 404


class TestExportarXlsx:
    def test_retorna_200(self, cliente, enxoval_id):
        r = cliente.get(f"/api/v1/enxovais/{enxoval_id}/exportar.xlsx")
        assert r.status_code == 200

    def test_content_type_xlsx(self, cliente, enxoval_id):
        r = cliente.get(f"/api/v1/enxovais/{enxoval_id}/exportar.xlsx")
        assert "spreadsheetml" in r.headers["content-type"]

    def test_content_disposition_attachment(self, cliente, enxoval_id):
        r = cliente.get(f"/api/v1/enxovais/{enxoval_id}/exportar.xlsx")
        assert "attachment" in r.headers["content-disposition"]
        assert ".xlsx" in r.headers["content-disposition"]

    def test_bytes_sao_zip_valido(self, cliente, enxoval_id):
        import openpyxl
        r = cliente.get(f"/api/v1/enxovais/{enxoval_id}/exportar.xlsx")
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        assert "Planilha" in wb.sheetnames
        assert "Roteiro" in wb.sheetnames

    def test_404_para_enxoval_inexistente(self, cliente):
        r = cliente.get("/api/v1/enxovais/00000000-0000-0000-0000-000000000000/exportar.xlsx")
        assert r.status_code == 404


class TestExportarPdf:
    def test_retorna_200(self, cliente, enxoval_id):
        r = cliente.get(f"/api/v1/enxovais/{enxoval_id}/exportar.pdf")
        assert r.status_code == 200

    def test_content_type_pdf(self, cliente, enxoval_id):
        r = cliente.get(f"/api/v1/enxovais/{enxoval_id}/exportar.pdf")
        assert "application/pdf" in r.headers["content-type"]

    def test_content_disposition_attachment(self, cliente, enxoval_id):
        r = cliente.get(f"/api/v1/enxovais/{enxoval_id}/exportar.pdf")
        assert "attachment" in r.headers["content-disposition"]
        assert ".pdf" in r.headers["content-disposition"]

    def test_bytes_comecam_com_assinatura_pdf(self, cliente, enxoval_id):
        r = cliente.get(f"/api/v1/enxovais/{enxoval_id}/exportar.pdf")
        assert r.content[:4] == b"%PDF"

    def test_404_para_enxoval_inexistente(self, cliente):
        r = cliente.get("/api/v1/enxovais/00000000-0000-0000-0000-000000000000/exportar.pdf")
        assert r.status_code == 404
