"""Gerador de XLSX da planilha do enxoval."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.rotas.schemas import EnxovalSaida

# Cores (ARGB sem #)
_COR_CABECALHO_FUNDO = "FF4F6B57"   # principal (verde-sálvia escuro)
_COR_CABECALHO_FONTE = "FFFFFFFF"   # branco
_COR_CATEGORIA_FUNDO = "FFEEF1EA"   # principal-suave
_COR_CATEGORIA_FONTE = "FF2F3A33"   # texto
_COR_FALTANDO_FUNDO  = "FFF7E4DF"   # alerta-fundo (salmão claro)

_ROTULO_PRIORIDADE = {
    "essencial": "Essencial",
    "util": "Útil",
    "pode_esperar": "Pode esperar",
}

_ROTULO_MOMENTO = {
    "atrasado": "⚠ Atrasado",
    "agora": "Agora",
    "proxima_fase": "Próxima fase",
    "futuro": "Mais para frente",
}


def _fill(cor: str) -> PatternFill:
    return PatternFill("solid", fgColor=cor)


def _ajustar_larguras(ws) -> None:
    """Ajusta a largura de cada coluna com base no conteúdo."""
    for col in ws.columns:
        largura = max((len(str(cell.value or "")) for cell in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(largura + 4, 40)


def _aba_planilha(wb: Workbook, saida: EnxovalSaida) -> None:
    ws = wb.create_sheet("Planilha")

    # ── cabeçalho ──────────────────────────────────────────────────────────
    cabecalhos = ["Item", "Tamanho", "Quantidade", "Prioridade",
                  "Momento", "Comprada", "Ganhada", "Já tinha", "Faltam"]
    for col, titulo in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=1, column=col, value=titulo)
        cell.font = Font(bold=True, color=_COR_CABECALHO_FONTE)
        cell.fill = _fill(_COR_CABECALHO_FUNDO)
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"
    linha_atual = 2

    # ── itens agrupados por categoria ──────────────────────────────────────
    categorias_ordenadas = sorted(saida.categorias, key=lambda c: c.ordem)

    for categoria in categorias_ordenadas:
        linhas_da_cat = [l for l in saida.linhas if l.categoria_slug == categoria.slug]
        if not linhas_da_cat:
            continue

        # Linha de subtítulo da categoria
        cell_cat = ws.cell(row=linha_atual, column=1, value=categoria.nome)
        cell_cat.font = Font(bold=True, color=_COR_CATEGORIA_FONTE)
        cell_cat.fill = _fill(_COR_CATEGORIA_FUNDO)
        ws.merge_cells(
            start_row=linha_atual, start_column=1,
            end_row=linha_atual, end_column=len(cabecalhos),
        )
        linha_atual += 1

        linhas_da_cat.sort(key=lambda l: (l.nome, l.tamanho or ""))
        for linha in linhas_da_cat:
            nome = linha.nome + (f" ({linha.rotulo_variante})" if linha.rotulo_variante else "")
            valores = [
                nome,
                linha.tamanho or "",
                linha.quantidade,
                _ROTULO_PRIORIDADE.get(linha.prioridade, linha.prioridade),
                _ROTULO_MOMENTO.get(linha.momento_compra, linha.momento_compra),
                linha.comprada,
                linha.ganhada,
                linha.ja_tinha,
                linha.faltam,
            ]
            for col, valor in enumerate(valores, start=1):
                cell = ws.cell(row=linha_atual, column=col, value=valor)
                if linha.faltam > 0:
                    cell.fill = _fill(_COR_FALTANDO_FUNDO)
            linha_atual += 1

    _ajustar_larguras(ws)


def _aba_roteiro(wb: Workbook, saida: EnxovalSaida) -> None:
    ws = wb.create_sheet("Roteiro")

    cabecalhos = ["Fase", "Início", "Fim", "Atual"]
    for col, titulo in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=1, column=col, value=titulo)
        cell.font = Font(bold=True, color=_COR_CABECALHO_FONTE)
        cell.fill = _fill(_COR_CABECALHO_FUNDO)

    for row, fase in enumerate(saida.roteiro, start=2):
        ws.cell(row=row, column=1, value=fase.nome)
        ws.cell(row=row, column=2, value=fase.inicio.strftime("%d/%m/%Y"))
        ws.cell(row=row, column=3, value=fase.fim.strftime("%d/%m/%Y"))
        atual_cell = ws.cell(row=row, column=4, value="✓" if fase.atual else "")
        if fase.atual:
            atual_cell.font = Font(bold=True, color=_COR_CABECALHO_FUNDO)

    _ajustar_larguras(ws)


def gerar_xlsx(saida: EnxovalSaida) -> bytes:
    """Devolve o XLSX completo como bytes."""
    wb = Workbook()
    # Remove a aba padrão criada pelo openpyxl
    del wb[wb.sheetnames[0]]

    _aba_planilha(wb, saida)
    _aba_roteiro(wb, saida)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
